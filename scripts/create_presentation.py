import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

# Colors
C_NAVY_DARK = RGBColor(11, 25, 44)      # #0B192C Deep Navy
C_NAVY_MED = RGBColor(26, 43, 76)       # #1A2B4C
C_BLUE_ACCENT = RGBColor(0, 102, 204)   # #0066CC
C_TEAL = RGBColor(2, 128, 144)          # #028090
C_AMBER = RGBColor(230, 92, 0)          # #E65C00 Safety Orange/Amber
C_GREEN = RGBColor(22, 101, 52)         # #166534 Success Green
C_GREEN_BG = RGBColor(240, 253, 244)    # #F0FDF4
C_RED = RGBColor(185, 28, 28)           # #B91C1C Danger Red
C_RED_BG = RGBColor(254, 242, 242)      # #FEF2F2
C_BLUE_BG = RGBColor(239, 246, 255)     # #EFF6FF
C_AMBER_BG = RGBColor(255, 247, 237)    # #FFF7ED
C_BG_SLATE = RGBColor(248, 250, 252)    # #F8FAFC
C_WHITE = RGBColor(255, 255, 255)
C_BORDER_GRAY = RGBColor(203, 213, 225) # #CBD5E1
C_TEXT_DARK = RGBColor(15, 23, 42)      # #0F172A
C_TEXT_MUTED = RGBColor(71, 85, 105)    # #475569
C_TEXT_LIGHT = RGBColor(241, 245, 249)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

def add_header_footer(slide, title_text, category_text, slide_num):
    # Top banner header
    top_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.3), Inches(12.133), Inches(0.8))
    tf = top_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p0 = tf.paragraphs[0]
    p0.text = category_text.upper()
    p0.font.size = Pt(11)
    p0.font.bold = True
    p0.font.color.rgb = C_BLUE_ACCENT
    p0.font.name = "Calibri"
    
    p1 = tf.add_paragraph()
    p1.text = title_text
    p1.font.size = Pt(26)
    p1.font.bold = True
    p1.font.color.rgb = C_NAVY_DARK
    p1.font.name = "Arial"
    
    # Bottom footer bar
    footer_bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.05), Inches(13.333), Inches(0.45))
    footer_bg.fill.solid()
    footer_bg.fill.fore_color.rgb = C_NAVY_DARK
    footer_bg.line.fill.background()
    
    footer_box = slide.shapes.add_textbox(Inches(0.6), Inches(7.05), Inches(12.133), Inches(0.45))
    ftf = footer_box.text_frame
    ftf.vertical_anchor = MSO_ANCHOR.MIDDLE
    ftf.margin_left = ftf.margin_top = ftf.margin_right = ftf.margin_bottom = 0
    fp = ftf.paragraphs[0]
    fp.text = "SIH26123 — Edge-AI Distributed AMR Fleet Coordination  |  Bharat Electronics Limited (BEL)"
    fp.font.size = Pt(12)
    fp.font.color.rgb = C_TEXT_LIGHT
    fp.font.name = "Calibri"
    
    # Slide number box
    num_box = slide.shapes.add_textbox(Inches(12.0), Inches(7.05), Inches(0.7), Inches(0.45))
    ntf = num_box.text_frame
    ntf.vertical_anchor = MSO_ANCHOR.MIDDLE
    np = ntf.paragraphs[0]
    np.text = str(slide_num)
    np.alignment = PP_ALIGN.RIGHT
    np.font.size = Pt(13)
    np.font.bold = True
    np.font.color.rgb = C_TEXT_LIGHT
    np.font.name = "Calibri"

print("Setup completed successfully.")
